"""استخراج متن از فایل‌های راهنما و تکه‌تکه کردن آن‌ها برای جستجو."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

from . import config
from .textutils import clean_text, normalize

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv"}


@dataclass
class Chunk:
    page: int
    ord: int
    text: str


def checksum(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()[:32]


# ------------------------------------------------------------------
# استخراج متن به تفکیک صفحه
# ------------------------------------------------------------------
def extract_pages(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext in {".txt", ".md", ".markdown", ".csv"}:
        return _extract_plain(path)
    raise ValueError(f"فرمت پشتیبانی‌نشده: {ext}")


def _extract_pdf(path: Path) -> list[str]:
    import pymupdf

    pages: list[str] = []
    with pymupdf.open(path) as doc:
        for page in doc:
            pages.append(clean_text(page.get_text("text")))
    return pages


def _extract_docx(path: Path) -> list[str]:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = clean_text("\n".join(parts))
    return _split_into_pseudo_pages(text)


def _extract_plain(path: Path) -> list[str]:
    text = clean_text(path.read_text(encoding="utf-8", errors="replace"))
    return _split_into_pseudo_pages(text)


def _split_into_pseudo_pages(text: str, size: int = 3000) -> list[str]:
    """فایل‌های بدون صفحه را به «صفحه»های مجازی تقسیم می‌کند تا ارجاع‌دهی ممکن شود."""
    if len(text) <= size:
        return [text]
    pages, buffer = [], []
    length = 0
    for para in text.split("\n\n"):
        if length + len(para) > size and buffer:
            pages.append("\n\n".join(buffer))
            buffer, length = [], 0
        buffer.append(para)
        length += len(para) + 2
    if buffer:
        pages.append("\n\n".join(buffer))
    return pages


# ------------------------------------------------------------------
# تکه‌بندی
# ------------------------------------------------------------------
def chunk_pages(
    pages: list[str],
    *,
    size: int | None = None,
    overlap: int | None = None,
) -> list[Chunk]:
    size = size or config.CHUNK_CHARS
    overlap = overlap or config.CHUNK_OVERLAP
    chunks: list[Chunk] = []
    order = 0

    for page_no, page_text in enumerate(pages, start=1):
        text = (page_text or "").strip()
        if len(normalize(text)) < 25:  # صفحه‌ی خالی یا فقط تصویر
            continue
        for piece in _split_text(text, size, overlap):
            chunks.append(Chunk(page=page_no, ord=order, text=piece))
            order += 1
    return chunks


def _split_text(text: str, size: int, overlap: int) -> list[str]:
    if len(text) <= size:
        return [text]

    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    pieces: list[str] = []
    buffer: list[str] = []
    length = 0

    for para in paragraphs:
        # پاراگرافِ خیلی بلند را با برش سخت می‌شکنیم
        while len(para) > size:
            head, para = para[:size], para[size - overlap :]
            pieces.append(head)
        if length + len(para) + 1 > size and buffer:
            pieces.append("\n".join(buffer))
            tail = "\n".join(buffer)[-overlap:]
            buffer = [tail] if tail else []
            length = len(tail)
        buffer.append(para)
        length += len(para) + 1

    if buffer:
        pieces.append("\n".join(buffer))
    return [p for p in pieces if len(p.strip()) > 30]
